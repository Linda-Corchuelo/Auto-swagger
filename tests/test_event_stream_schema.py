"""
Pruebas del soporte de esquemas para respuestas `text/event-stream` (SSE).

Contexto
--------
El endpoint `GET /events/stream` declara su respuesta 200 como:

    "content": { "text/event-stream": { "schema": { "type": "string",
                                                     "properties": [],
                                                     "example": "{...}" } } }

Como el schema no tiene `properties`, `add_schema_properties_table`
retornaba sin dibujar nada y la tabla de "Respuesta" salía vacía.

La opción (c) hace que el parser reconstruya el schema a partir del
`example` cuando el content type es `text/event-stream`.

Ejecución
---------
    python tests/test_event_stream_schema.py
"""

import json
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from swagger.parser import (  # noqa: E402
    build_schema_from_example,
    parse_content_schema,
    parse_endpoints,
)

# Ejemplo real publicado por el microservicio de events (SSE)
SSE_EXAMPLE = (
    '{ "id": "03a6c6e8-452d-46c1-9b85-0753901608b6",'
    ' "payload": {"title": "Prueba",'
    ' "message": "notificacion de pruebas",'
    ' "link": "https://google.com",'
    ' "type": "success"},'
    ' "userName": "124",'
    ' "read": false,'
    ' "timestamp": "2026-10-08T14:30:43.648Z"}'
)


def _names(properties):
    return [prop.name for prop in properties]


# ============================================================
# 1. build_schema_from_example
# ============================================================

def test_build_schema_desde_example_sse():
    schema = build_schema_from_example("AnonymousSchema", SSE_EXAMPLE)

    assert schema is not None, "No se construyó el schema desde el example"
    assert schema.type == "object"
    assert _names(schema.properties) == [
        "id",
        "payload",
        "userName",
        "read",
        "timestamp",
    ], f"Propiedades inesperadas: {_names(schema.properties)}"

    # El example original (string) se conserva intacto
    assert schema.example == SSE_EXAMPLE

    # Tipos inferidos correctamente
    types = {prop.name: prop.type for prop in schema.properties}
    assert types["id"] == "string"
    assert types["read"] == "boolean"
    assert types["timestamp"] == "string"

    # Anidamiento de `payload`
    payload = next(p for p in schema.properties if p.name == "payload")
    assert payload.type == "object"
    assert _names(payload.properties) == [
        "title",
        "message",
        "link",
        "type",
    ]


def test_build_schema_acepta_dict():
    schema = build_schema_from_example(
        "Resp",
        {"code": "ES01", "count": 3, "ok": True, "items": [1, 2]},
    )

    assert schema is not None
    assert schema.type == "object"
    types = {prop.name: prop.type for prop in schema.properties}
    assert types == {
        "code": "string",
        "count": "integer",
        "ok": "boolean",
        "items": "array",
    }

    items = next(p for p in schema.properties if p.name == "items")
    assert items.items is not None
    assert items.items.type == "integer"


def test_build_schema_retorna_none_en_casos_invalidos():
    # Sin example
    assert build_schema_from_example("S", None) is None

    # String que no es JSON
    assert build_schema_from_example("S", "hola mundo") is None

    # JSON escalar (sin estructura de tabla)
    assert build_schema_from_example("S", "42") is None
    assert build_schema_from_example("S", '"texto"') is None

    # Objeto vacío
    assert build_schema_from_example("S", "{}") is None
    assert build_schema_from_example("S", {}) is None

    # Lista (raíz array): fuera de alcance de la tabla de propiedades
    assert build_schema_from_example("S", "[1,2,3]") is None


# ============================================================
# 2. parse_content_schema
# ============================================================

def _sse_media_type(example=SSE_EXAMPLE):
    return {
        "text/event-stream": {
            "schema": {
                "type": "string",
                "example": example,
            }
        }
    }


def test_parse_content_schema_reconstruye_event_stream():
    result = parse_content_schema(_sse_media_type(), document={})
    schema = result["text/event-stream"]

    assert schema.type == "object"
    assert len(schema.properties) == 5, "No se reconstruyeron las propiedades"
    assert schema.example == SSE_EXAMPLE


def test_parse_content_schema_json_regular_sin_cambios():
    """Regresión: application/json NO debe sufrir la reconstrucción."""
    content = {
        "application/json": {
            "schema": {
                "$ref": "#/components/schemas/ApiResponse",
            }
        }
    }
    document = {
        "components": {
            "schemas": {
                "ApiResponse": {
                    "type": "object",
                    "properties": {"statusCode": {"type": "integer"}},
                }
            }
        }
    }

    result = parse_content_schema(content, document)
    schema = result["application/json"]

    assert schema.type == "object"
    assert _names(schema.properties) == ["statusCode"], (
        "El schema JSON fue alterado por la nueva lógica"
    )


def test_parse_content_schema_event_stream_sin_example_valido():
    """Si el example no es JSON, se conserva el schema original."""
    result = parse_content_schema(
        _sse_media_type(example="no es json"), document={}
    )
    schema = result["text/event-stream"]

    assert schema.type == "string"
    assert schema.properties == []
    assert schema.example == "no es json"


def test_parse_content_schema_event_stream_con_properties_no_toca_nada():
    """Si el SSE ya trae properties declaradas, no se reconstruye."""
    content = {
        "text/event-stream": {
            "schema": {
                "type": "object",
                "properties": {"id": {"type": "string"}},
            }
        }
    }
    result = parse_content_schema(content, document={})
    schema = result["text/event-stream"]

    assert _names(schema.properties) == ["id"]


# ============================================================
# 3. Integración: endpoint completo + tabla en Word
# ============================================================

def _sse_openapi_document():
    return {
        "openapi": "3.0.0",
        "info": {"title": "Rummi Events API", "version": "1.0.0"},
        "paths": {
            "/events/stream": {
                "get": {
                    "summary": "Suscripción a notificaciones del usuario",
                    "description": (
                        "Mantiene una conexión Server-Sent Events abierta, "
                        "No usa cifrado."
                    ),
                    "parameters": [
                        {
                            "name": "userId",
                            "in": "query",
                            "required": True,
                            "description": "Identificador del usuario.",
                            "schema": {"type": "string"},
                        }
                    ],
                    "responses": {
                        "200": {
                            "description": "Evento SSE.",
                            "content": _sse_media_type(),
                        },
                        "400": {
                            "description": "userId obligatorio.",
                            "content": {
                                "application/json": {
                                    "schema": {
                                        "type": "object",
                                        "required": ["error"],
                                        "properties": {
                                            "error": {"type": "string"}
                                        },
                                    }
                                }
                            },
                        },
                    },
                }
            }
        },
    }


def test_parse_endpoints_stream_con_respuesta_con_propiedades():
    endpoints = parse_endpoints(_sse_openapi_document())
    assert len(endpoints) == 1

    endpoint = endpoints[0]
    resp_200 = endpoint.responses[0]
    schema = resp_200.content["text/event-stream"]

    assert schema.type == "object"
    assert len(schema.properties) == 5
    assert schema.example == SSE_EXAMPLE


def test_tabla_respuesta_stream_con_filas():
    """La tabla de la respuesta ahora debe renderizar filas."""
    from docx import Document
    from generator.word_generator import (
        add_schema_properties_table,
        collect_schema_rows,
    )

    endpoint = parse_endpoints(_sse_openapi_document())[0]
    resp_schema = endpoint.responses[0].content["text/event-stream"]

    # 1) Filas planas que usaría la tabla
    rows = collect_schema_rows(resp_schema.properties)
    names = [row["name"] for row in rows]
    assert "id" in names and "payload" in names and "timestamp" in names, (
        f"Faltan filas: {names}"
    )
    assert "payload.title" in names, "No se expandió el anidamiento"
    assert all(row["description"] for row in rows), (
        "Hay filas sin descripción"
    )

    # 2) Tabla real dentro del documento
    document = Document()
    add_schema_properties_table(document, resp_schema)

    assert len(document.tables) == 1, "No se generó la tabla de respuesta"
    table = document.tables[0]
    assert len(table.rows) == 1 + len(rows), (
        f"Filas esperadas {1 + len(rows)}, obtenidas {len(table.rows)}"
    )
    assert [
        cell.text for cell in table.rows[0].cells
    ] == ["Nombre", "Tipo", "Tamaño", "Obligatorio", "Descripción"]


def test_ejemplo_json_de_salida_se_mantiene():
    """El bloque de ejemplo debe seguir mostrando el example original."""
    from generator.word_generator import generate_sample_json

    endpoint = parse_endpoints(_sse_openapi_document())[0]
    resp_schema = endpoint.responses[0].content["text/event-stream"]

    sample = generate_sample_json(resp_schema)
    assert sample == SSE_EXAMPLE, (
        "El ejemplo de salida cambió de formato"
    )


# ============================================================
# Runner
# ============================================================

def main():
    tests = [
        value
        for name, value in sorted(globals().items())
        if name.startswith("test_") and callable(value)
    ]

    failures = []

    for test in tests:
        try:
            test()
        except Exception as error:  # noqa: BLE001
            failures.append((test.__name__, error))
            print(f"[FALLO] {test.__name__}: {type(error).__name__}: {error}")
        else:
            print(f"[ OK ] {test.__name__}")

    print("-" * 60)
    print(f"Pruebas: {len(tests)} | OK: {len(tests) - len(failures)} | "
          f"Fallos: {len(failures)}")

    return 1 if failures else 0


if __name__ == "__main__":
    sys.exit(main())
